"""AppTest coverage for uvalu/pages_/settings.py."""
import io

import pandas as pd
import pytest
import streamlit as st
from streamlit.testing.v1 import AppTest

import auth
import portfolio
import settings
from uvalu import oauth
from uvalu.pages_ import settings as settings_page
from tests.conftest import TEST_EMAIL, fake_cached_fn, USER_SETUP_SRC


class FakeUser(dict):
    """Minimal stand-in for streamlit.user_info.UserInfoProxy — see
    tests/test_oauth.py's identical class for why this suffices."""
    def __init__(self, is_logged_in: bool, **claims):
        super().__init__(**claims)
        self.is_logged_in = is_logged_in


def _run(monkeypatch, role="Analyst", jwt_sid=None) -> AppTest:
    # AppTest.from_function re-executes a function's SOURCE TEXT as a fresh
    # script (no closure over enclosing variables — see tests/test_pages_admin.py's
    # identical note), so `role` can't be passed as a real Python value into a
    # from_function script; AppTest.from_string takes the already-interpolated
    # text directly instead.
    monkeypatch.setattr(settings_page, "_load_all_screener_data", fake_cached_fn(None))

    _sid_line = f'st.session_state["jwt_sid"] = {jwt_sid!r}\n' if jwt_sid else ""
    script_src = USER_SETUP_SRC + f"""
import streamlit as st
from uvalu.pages_ import settings as settings_page
st.session_state["user_email"] = "test@example.com"
st.session_state["user_role"] = {role!r}
{_sid_line}settings_page.render()
"""
    at = AppTest.from_string(script_src, default_timeout=60)
    at.run()
    assert not at.exception, [str(e.value) for e in at.exception]
    return at


def test_renders_without_exceptions(isolated_data, monkeypatch):
    _run(monkeypatch)


def test_model_settings_in_force_list_for_every_user(isolated_data, monkeypatch):
    """G14: a read-only list of the settings in force, the same for everyone."""
    at = _run(monkeypatch)
    html = "".join(m.value for m in at.markdown)
    assert "Model settings in force" in html and "Same for every user · model M1" in html
    assert "&gt; 500%" in html          # quality screen: debt / equity
    assert "≥ 70" in html and "≥ 10%" in html
    assert "3% + β × 5%" in html and "Terminal dividend growth" in html
    assert "Margin of safety 24 · Risk 22 · Quality 24 · Momentum 15 · Dividend 15." in html
    assert "Changes are logged." in html and "announced" not in html


def test_changing_max_debt_equity_slider_persists_and_reruns(isolated_data, monkeypatch):
    # Screening & veto rules are admin-only (see the non-admin tests below).
    at = _run(monkeypatch, role="Admin")
    slider = at.slider(key="scr_max_de")
    slider.set_value(300)
    at.run()
    assert not at.exception, [str(e.value) for e in at.exception]
    assert settings.load_shared_settings()["max_debt_equity"] == 300.0


def test_changing_benchmark_toggle_persists(isolated_data, monkeypatch):
    at = _run(monkeypatch, role="Admin")
    toggle = at.toggle(key="scr_stoxx")
    toggle.set_value(True)
    at.run()
    assert not at.exception, [str(e.value) for e in at.exception]
    assert settings.load_shared_settings()["benchmark_stoxx"] is True


@pytest.mark.parametrize("role", ["Analyst", "Viewer"])
def test_screening_sliders_only_for_admins(isolated_data, monkeypatch, role):
    # Shared, all-user settings: a non-admin sees the list, never a control
    # that could change what every other user's signals are based on.
    at = _run(monkeypatch, role=role)
    keys = {w.key for w in at.slider} | {w.key for w in at.segmented_control}
    assert not keys & {"scr_max_de", "scr_max_payout", "scr_min_mos", "scr_buy_thr", "scr_style"}
    assert at.toggle(key="scr_stoxx").disabled is True


def test_changing_screening_style_persists_and_reruns(isolated_data, monkeypatch):
    at = _run(monkeypatch, role="Admin")
    at.segmented_control(key="scr_style").set_value("Income")
    at.run()
    assert not at.exception, [str(e.value) for e in at.exception]
    assert settings.load_shared_settings()["screen_style"] == "income"


def test_export_is_offered_when_portfolio_empty(isolated_data, monkeypatch):
    # An empty account's export is the template for a bulk import.
    at = _run(monkeypatch)
    assert [b.label for b in at.download_button] == ["Download backup.xlsx", "Export"]


def test_export_shows_download_button_when_portfolio_has_data(isolated_data, monkeypatch):
    portfolio.save_portfolio(pd.DataFrame([{"ticker": "AAA.BR", "shares": 10}]))
    at = _run(monkeypatch)
    assert [b.label for b in at.download_button] == ["Download backup.xlsx", "Export"]


def _build_test_workbook() -> bytes:
    """Rows added to an empty export: one position, one sold position."""
    import openpyxl
    import workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Positions")
    ws.append(list(workbook.COLUMNS["positions"]))
    ws.append(["Test Corp", "TESTX.BR", 10, 1000.0, None, "2024-01-01"])
    ws = wb.create_sheet("Sold")
    ws.append(list(workbook.COLUMNS["sold"]))
    ws.append(["Sold Corp", "SOLDX.BR", 5, 500.0, 600.0, None, None, "2023-01-01", "2023-06-01"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _no_portfolio() -> bool:
    pf = portfolio.load_portfolio()
    return pf is None or pf.empty


def _upload_workbook(at, data):
    at.file_uploader(key="imp_portfolio").upload(
        "portfolio.xlsx", data,
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    at.run()


def test_import_adds_after_confirmation(isolated_data, monkeypatch):
    # Nothing is written until "Import". The uploader gets a new key
    # afterwards — the old one kept its file and re-imported it on every
    # rerun.
    at = _run(monkeypatch)
    _upload_workbook(at, _build_test_workbook())
    assert not at.exception, [str(e.value) for e in at.exception]
    assert _no_portfolio()                                      # nothing written yet
    assert "This file holds" in "".join(m.value for m in at.markdown)
    at.button(key="imp_add").click().run()
    assert not at.exception, [str(e.value) for e in at.exception]
    assert portfolio.load_portfolio().iloc[0]["ticker"] == "TESTX.BR"
    assert portfolio.load_sold().iloc[0]["ticker"] == "SOLDX.BR"
    # A fresh, empty uploader replaced the one holding the file: no re-import
    # loop. (AppTest keeps the pre-rerun buttons in its tree, so assert on
    # the uploader rather than the confirm button's absence.)
    assert [(u.key, bool(u.value)) for u in at.file_uploader] == [("imp_portfolio_1", False)]


def test_import_replace_asks_for_the_red_confirmation(isolated_data, monkeypatch):
    portfolio.save_portfolio(pd.DataFrame([{"ticker": "OLD.BR", "shares": 1, "purchase_value": 10.0,
                                            "date_in": "2020-01-01T00:00:00"}]))
    at = _run(monkeypatch)
    _upload_workbook(at, _build_test_workbook())
    at.radio(key="imp_mode_0").set_value("replace").run()
    assert "All your portfolio data is deleted first" in "".join(m.value for m in at.markdown)
    at.button(key="imp_confirm").click().run()
    assert not at.exception, [str(e.value) for e in at.exception]
    assert list(portfolio.load_portfolio()["ticker"]) == ["TESTX.BR"]


def test_import_cancel_keeps_the_current_portfolio(isolated_data, monkeypatch):
    at = _run(monkeypatch)
    _upload_workbook(at, _build_test_workbook())
    at.button(key="imp_cancel").click().run()
    assert not at.exception, [str(e.value) for e in at.exception]
    assert _no_portfolio()
    assert [(u.key, bool(u.value)) for u in at.file_uploader] == [("imp_portfolio_1", False)]


def test_import_lists_row_problems(isolated_data, monkeypatch):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ticker", "shares", "purchase_value", "date_in"])
    ws.append(["TESTX.BR", "ten", 1000, "2024-01-01"])
    buf = io.BytesIO()
    wb.save(buf)
    at = _run(monkeypatch)
    _upload_workbook(at, buf.getvalue())
    assert not at.exception, [str(e.value) for e in at.exception]
    err = "".join(e.value for e in at.error)
    assert "row 2" in err and "not a number" in err
    assert not [b for b in at.button if b.key in ("imp_add", "imp_confirm")]


def test_targets_section_renders(isolated_data, monkeypatch):
    at = _run(monkeypatch)
    html = "".join(m.value for m in at.markdown)
    assert "Target allocation" in html
    assert at.text_area(key="tgt_sectors").value == ""      # nothing saved yet
    assert at.number_input(key="tgt_hhi").value is None         # empty = default bands


def test_save_target_allocation_persists(isolated_data, monkeypatch):
    import streamlit as st
    monkeypatch.setattr(st, "rerun", lambda *a, **k: None)   # save branch reruns

    at = _run(monkeypatch)
    at.text_area(key="tgt_sectors").set_value("Technology 25\nHealthcare 15")
    at.text_area(key="tgt_tickers").set_value("AAA.BR 10")
    at.number_input(key="tgt_hhi").set_value(0.15)
    at.run()
    at.button(key="tgt_save").click()
    at.run()
    assert not at.exception, [str(e.value) for e in at.exception]

    t = portfolio.load_targets()
    assert t["sectors"] == {"Technology": 0.25, "Healthcare": 0.15}
    assert t["tickers"] == {"AAA.BR": 0.10}
    assert t["hhi_max"] == 0.15


def test_existing_targets_prefill_the_editor(isolated_data, monkeypatch):
    portfolio.save_targets({"sectors": {"Technology": 0.30}, "hhi_max": 0.12})
    at = _run(monkeypatch)
    assert "Technology 30" in at.text_area(key="tgt_sectors").value
    assert at.number_input(key="tgt_hhi").value == 0.12


def test_withholding_rates_save_only_with_the_button(isolated_data, monkeypatch):
    at = _run(monkeypatch)
    at.number_input(key="disp_wh_brussels").set_value(30.0).run()
    assert settings.load_settings(TEST_EMAIL).get("dividend_withholding", {}).get("brussels", 0.0) == 0.0
    at.button(key="wh_save").click().run()
    assert not at.exception, [str(e.value) for e in at.exception]
    assert settings.load_settings(TEST_EMAIL)["dividend_withholding"]["brussels"] == 30.0


def test_section_nav_links_every_card_anchor(isolated_data, monkeypatch):
    at = _run(monkeypatch)
    html = "".join(m.value for m in at.markdown)
    for anchor in ("set-language", "set-account", "set-display", "set-portfolio", "set-withholding",
                   "set-import", "set-model", "set-legal"):
        assert f'href="#{anchor}"' in html and f'id="{anchor}"' in html, anchor
    assert "set-data" not in html and "Data &amp; alerts" not in html


def test_settings_order_matches_reference():
    """Rail order of the design reference: … Target allocation, Dividend
    withholding, Import & export …, Privacy, Plan, Legal; no Data & alerts."""
    from uvalu.pages_.settings import _SECTIONS
    assert [a for a, *_ in _SECTIONS] == ["set-language", "set-account", "set-display", "set-portfolio",
                                         "set-withholding", "set-import", "set-model", "set-privacy",
                                         "set-plan", "set-legal"]


def test_parse_targets_text_drops_malformed_and_out_of_range():
    from uvalu.pages_.settings import _parse_targets_text
    parsed = _parse_targets_text("Technology 25\nBad\nHuge 250\nConsumer Cyclical 12%\n  \nZero 0")
    assert parsed == {"Technology": 0.25, "Consumer Cyclical": 0.12}


def test_account_footer_comes_last_with_email_and_free_beta(isolated_data, monkeypatch):
    """ST-05 / G27: after the last card; "Free beta" instead of the role."""
    at = _run(monkeypatch)
    html = "".join(m.value for m in at.markdown)
    assert TEST_EMAIL.replace("@", "&#64;") + " · Free beta" in html
    assert "Analyst" not in html
    assert html.index("uv-set-account-footer") > html.index('id="set-legal"')
    assert "Sign out" in html


# ── logging (logkit Phase 2) ─────────────────────────────────────────────

import logging as _logging  # noqa: E402

from uvalu import logkit  # noqa: E402


def _config_changes(caplog, key=None):
    return [r for r in caplog.records
            if getattr(r, "event", None) == "config.change"
            and (key is None or getattr(r, "key", None) == key)]


class TestSettingsLogging:
    def test_shared_settings_change_logs_per_key_with_old_and_new(self, isolated_data, caplog):
        caplog.set_level(_logging.DEBUG, logger="uvalu")
        base = settings.load_shared_settings()
        settings.save_shared_settings({**base, "max_debt_equity": 400.0, "screen_style": "value"})
        changed = {r.key: (r.old, r.new, r.scope) for r in _config_changes(caplog)}
        assert changed["max_debt_equity"] == (500.0, 400.0, "shared")
        assert changed["screen_style"][1] == "value"
        assert "buy_threshold" not in changed          # unchanged keys don't log

    def test_user_settings_change_scope_is_user_hash(self, isolated_data, caplog):
        caplog.set_level(_logging.DEBUG, logger="uvalu")
        settings.save_settings({**settings.load_settings(TEST_EMAIL), "density": "compact"}, TEST_EMAIL)
        (rec,) = _config_changes(caplog, "density")
        assert rec.new == "compact"
        assert rec.scope == f"user:{logkit.user_hash(TEST_EMAIL)}"

    def test_no_op_save_logs_nothing(self, isolated_data, caplog):
        caplog.set_level(_logging.DEBUG, logger="uvalu")
        current = settings.load_shared_settings()
        settings.save_shared_settings(dict(current))
        assert _config_changes(caplog) == []

    def test_corrupt_shared_settings_file_logs_and_falls_back(self, isolated_data, caplog):
        caplog.set_level(_logging.DEBUG, logger="uvalu")
        settings._SHARED_FILE.parent.mkdir(parents=True, exist_ok=True)
        settings._SHARED_FILE.write_bytes(b"not a valid encrypted payload")
        got = settings.load_shared_settings()
        assert got["max_debt_equity"] == 500.0           # defaults, behaviour unchanged
        reads = [r for r in caplog.records if getattr(r, "event", None) == "storage.read_failed"]
        assert reads and reads[0].file == "shared.json" and reads[0].exc_info is not None


# ── Security card ─────────────────────────────────────────────────────────

class TestSecurityCard:
    def test_shows_password_last_changed(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        at = _run(monkeypatch)
        html = "".join(m.value for m in at.markdown)
        assert "Last changed" in html

    def test_no_account_shows_set_a_password(self, isolated_data, monkeypatch):
        # TEST_EMAIL isn't registered in auth's user store at all — only
        # portfolio.set_user() was called (see USER_SETUP_SRC) — has_password()
        # is False for a nonexistent account, same UI as a real provider-only
        # account with no password_hash yet.
        at = _run(monkeypatch)
        html = "".join(m.value for m in at.markdown)
        assert "NOT SET" in html
        assert any(b.label == "Set a password" for b in at.button)

    def test_change_button_opens_dialog(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        at = _run(monkeypatch)
        open_btn = [b for b in at.button if b.label == "Change"][0]
        open_btn.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert any(w.key == "set_pw_current" for w in at.text_input)


def _run_dlg_change_password(monkeypatch, email=TEST_EMAIL) -> AppTest:
    # Same one-shot-gate limitation as admin.py's _dlg_invite (see the long
    # comment above TestInviteUserButtonWiring in test_pages_admin.py) —
    # call the @st.dialog function directly and unconditionally instead of
    # driving it through the "Change" button click.
    script = f"""
from uvalu.pages_.settings import _dlg_change_password
_dlg_change_password({email!r})
"""
    at = AppTest.from_string(script, default_timeout=60)
    at.run()
    assert not at.exception, [str(e.value) for e in at.exception]
    return at


class TestDlgChangePassword:
    def test_correct_current_password_changes_it(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        at = _run_dlg_change_password(monkeypatch)
        at.text_input(key="set_pw_current").set_value("password123")
        at.text_input(key="set_pw_new").set_value("newpassword456")
        at.text_input(key="set_pw_confirm").set_value("newpassword456")
        submit = [b for b in at.button if b.label == "Change password"][0]
        submit.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert len(at.success) == 1
        ok, _ = auth.login(TEST_EMAIL, "newpassword456")
        assert ok

    def test_wrong_current_password_shows_error(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        at = _run_dlg_change_password(monkeypatch)
        at.text_input(key="set_pw_current").set_value("wrong-password")
        at.text_input(key="set_pw_new").set_value("newpassword456")
        at.text_input(key="set_pw_confirm").set_value("newpassword456")
        submit = [b for b in at.button if b.label == "Change password"][0]
        submit.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert "incorrect" in "".join(e.value for e in at.error).lower()

    def test_mismatched_confirmation_shows_error(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        at = _run_dlg_change_password(monkeypatch)
        at.text_input(key="set_pw_current").set_value("password123")
        at.text_input(key="set_pw_new").set_value("newpassword456")
        at.text_input(key="set_pw_confirm").set_value("different789")
        submit = [b for b in at.button if b.label == "Change password"][0]
        submit.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert "match" in "".join(e.value for e in at.error).lower()

    def test_new_password_shows_live_strength_caption(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        at = _run_dlg_change_password(monkeypatch)
        at.text_input(key="set_pw_new").set_value("weak")
        at.run()
        html = "".join(m.value for m in at.markdown)
        assert "Password strength: Weak" in html

        at.text_input(key="set_pw_new").set_value("Str0ng!-Passphrase-99")
        at.run()
        html = "".join(m.value for m in at.markdown)
        assert "Password strength: Strong" in html


def _run_dlg_set_password(monkeypatch, email=TEST_EMAIL) -> AppTest:
    script = f"""
from uvalu.pages_.settings import _dlg_set_password
_dlg_set_password({email!r})
"""
    at = AppTest.from_string(script, default_timeout=60)
    at.run()
    assert not at.exception, [str(e.value) for e in at.exception]
    return at


class TestDlgSetPassword:
    def test_new_password_shows_live_strength_caption(self, isolated_data, monkeypatch):
        at = _run_dlg_set_password(monkeypatch)
        at.text_input(key="set_pw2_new").set_value("weak")
        at.run()
        html = "".join(m.value for m in at.markdown)
        assert "Password strength: Weak" in html


# ── Active sessions card ──────────────────────────────────────────────────

class TestActiveSessions:
    def test_shows_current_and_other_sessions(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        _, tok1 = auth.login(TEST_EMAIL, "password123",
                             user_agent="Mozilla/5.0 (Macintosh) AppleWebKit Chrome/120.0")
        _, _, sid1 = auth.verify_token(tok1)
        auth.login(TEST_EMAIL, "password123",
                  user_agent="Mozilla/5.0 (Windows NT 10.0) AppleWebKit Firefox/120.0")
        at = _run(monkeypatch, jwt_sid=sid1)
        html = "".join(m.value for m in at.markdown)
        assert "This session" in html and "Signed in " in html
        assert "Chrome on macOS" in html
        assert "Firefox on Windows" in html
        assert any(b.label == "Sign out" for b in at.button)

    def test_current_session_has_no_sign_out_button(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        _, tok1 = auth.login(TEST_EMAIL, "password123")
        _, _, sid1 = auth.verify_token(tok1)
        at = _run(monkeypatch, jwt_sid=sid1)
        assert not any(b.key == f"set_session_signout_{sid1}" for b in at.button)

    def test_sign_out_one_session_revokes_only_that_one(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        _, tok1 = auth.login(TEST_EMAIL, "password123")
        _, _, sid1 = auth.verify_token(tok1)
        _, tok2 = auth.login(TEST_EMAIL, "password123")
        _, _, sid2 = auth.verify_token(tok2)
        at = _run(monkeypatch, jwt_sid=sid1)
        signout_btn = [b for b in at.button if b.key == f"set_session_signout_{sid2}"][0]
        signout_btn.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert auth.is_session_active(TEST_EMAIL, sid1)
        assert not auth.is_session_active(TEST_EMAIL, sid2)

    def test_sign_out_everywhere_else_keeps_current_only(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        _, tok1 = auth.login(TEST_EMAIL, "password123")
        _, _, sid1 = auth.verify_token(tok1)
        auth.login(TEST_EMAIL, "password123")
        auth.login(TEST_EMAIL, "password123")
        at = _run(monkeypatch, jwt_sid=sid1)
        revoke_btn = [b for b in at.button if b.label == "Sign out everywhere else"][0]
        revoke_btn.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert auth.is_session_active(TEST_EMAIL, sid1)
        assert len(auth.list_sessions(TEST_EMAIL)) == 1

    def test_single_session_hides_sign_out_everywhere_button(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        _, tok1 = auth.login(TEST_EMAIL, "password123")
        _, _, sid1 = auth.verify_token(tok1)
        at = _run(monkeypatch, jwt_sid=sid1)
        assert not any(b.label == "Sign out everywhere else" for b in at.button)


# ── Linked accounts ───────────────────────────────────────────────────────

class TestLinkedAccounts:
    def test_unlinked_configured_provider_shows_connect(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        monkeypatch.setattr(oauth, "_auth_secrets", lambda: {
            "google": {"client_id": "x", "client_secret": "y",
                      "server_metadata_url": "https://accounts.google.com/.well-known/openid-configuration"}})
        at = _run(monkeypatch)
        assert any(b.label == "Connect" for b in at.button)

    def test_google_row_reads_not_in_the_beta(self, isolated_data, monkeypatch):
        """ST-03: one "Linked account · Google" row while Google is off."""
        auth.register(TEST_EMAIL, "password123")
        at = _run(monkeypatch)
        html = "".join(m.value for m in at.markdown)
        assert "Linked account · Google" in html and "Not in the beta" in html
        assert "Sign-in with Google is not available in the beta." in html

    def test_sessions_sit_inside_the_security_card(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        auth.login(TEST_EMAIL, "password123", user_agent="Mozilla/5.0 (Macintosh) AppleWebKit Chrome/120.0")
        at = _run(monkeypatch)
        assert not at.exception, [str(e.value) for e in at.exception]
        assert not any("set_card_sessions" in (c.key or "") for c in at.get("flex_container"))
        html = "".join(m.value for m in at.markdown)
        assert html.index("Active sessions") > html.index("Two-factor authentication")

    def test_unconfigured_provider_shows_unavailable(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        at = _run(monkeypatch)
        html = "".join(m.value for m in at.markdown)
        assert "not available in the beta" in html
        assert not any(b.label == "Connect" for b in at.button)

    def test_linked_provider_shows_connected_and_disconnect(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        auth.link_identity(TEST_EMAIL, "https://accounts.google.com", "sub-123", "test@example.com")
        at = _run(monkeypatch)
        html = "".join(m.value for m in at.markdown)
        assert "CONNECTED" in html
        assert any(b.label == "Disconnect" for b in at.button)

    def test_disconnect_unlinks_provider(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        auth.link_identity(TEST_EMAIL, "https://accounts.google.com", "sub-123", "test@example.com")
        at = _run(monkeypatch)
        disconnect_btn = [b for b in at.button if b.label == "Disconnect"][0]
        disconnect_btn.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert auth.list_linked_identities(TEST_EMAIL) == []

    def test_disconnecting_only_method_is_refused(self, isolated_data, monkeypatch):
        # Password-less account with exactly one linked identity -- refused
        # by auth.unlink_identity(), surfaced here as a toast rather than
        # silently vanishing.
        _, _, token = auth.invite_user(TEST_EMAIL)
        auth.accept_invite_with_oauth(token, "https://accounts.google.com", "sub-123", TEST_EMAIL)
        at = _run(monkeypatch)
        disconnect_btn = [b for b in at.button if b.label == "Disconnect"][0]
        disconnect_btn.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert len(auth.list_linked_identities(TEST_EMAIL)) == 1

    def test_connect_click_starts_login(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        monkeypatch.setattr(oauth, "_auth_secrets", lambda: {
            "google": {"client_id": "x", "client_secret": "y",
                      "server_metadata_url": "https://accounts.google.com/.well-known/openid-configuration"}})
        started = {"provider": None}
        monkeypatch.setattr(oauth, "start_login", lambda pid: started.__setitem__("provider", pid))
        at = _run(monkeypatch)
        connect_btn = [b for b in at.button if b.label == "Connect"][0]
        connect_btn.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert started["provider"] == "google"


class TestOAuthSelfLinking:
    def test_completed_identity_links_to_current_account(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        monkeypatch.setattr(st, "user", FakeUser(
            True, iss="https://accounts.google.com", sub="sub-123", email=TEST_EMAIL))
        at = _run(monkeypatch)
        assert not at.exception, [str(e.value) for e in at.exception]
        linked = auth.list_linked_identities(TEST_EMAIL)
        assert len(linked) == 1
        assert linked[0]["subject"] == "sub-123"

    def test_identity_already_linked_elsewhere_is_not_stolen(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        auth.register("other@example.com", "password123456")
        auth.link_identity("other@example.com", "https://accounts.google.com", "sub-123", "other@example.com")
        monkeypatch.setattr(st, "user", FakeUser(
            True, iss="https://accounts.google.com", sub="sub-123", email=TEST_EMAIL))
        at = _run(monkeypatch)
        assert not at.exception, [str(e.value) for e in at.exception]
        assert auth.list_linked_identities(TEST_EMAIL) == []
        assert len(auth.list_linked_identities("other@example.com")) == 1


# ── Two-factor authentication ─────────────────────────────────────────────

def _enable_totp(email=TEST_EMAIL):
    import pyotp
    secret, _ = auth.begin_totp_enrollment(email)
    auth.confirm_totp_enrollment(email, pyotp.TOTP(secret).now())


class TestTotpToggleRow:
    def test_provider_only_account_shows_not_applicable(self, isolated_data, monkeypatch):
        # TEST_EMAIL isn't registered at all -- has_password() is False, same
        # as a real provider-only account.
        at = _run(monkeypatch)
        html = "".join(m.value for m in at.markdown)
        assert "Not applicable" in html
        assert not any(w.key == "set_totp_toggle" for w in at.toggle)

    def test_toggle_off_by_default(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        at = _run(monkeypatch)
        assert not at.toggle(key="set_totp_toggle").value

    def test_turning_on_opens_enrollment_dialog(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        at = _run(monkeypatch)
        at.toggle(key="set_totp_toggle").set_value(True)
        at.run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert any(w.key == "totp_enroll_code" for w in at.text_input)
        assert any(b.label == "Confirm and enable" for b in at.button)

    def test_turning_off_disables_totp(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        _enable_totp()
        at = _run(monkeypatch)
        at.toggle(key="set_totp_toggle").set_value(False)
        at.run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert not auth.is_totp_enabled(TEST_EMAIL)

    def test_enabled_shows_backup_codes_and_trusted_devices_rows(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        _enable_totp()
        at = _run(monkeypatch)
        html = "".join(m.value for m in at.markdown)
        assert "Backup codes" in html
        assert "Trusted devices" in html

    def test_disabled_hides_backup_codes_and_trusted_devices_rows(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        at = _run(monkeypatch)
        html = "".join(m.value for m in at.markdown)
        assert "Backup codes" not in html
        assert "Trusted devices" not in html


def _run_dlg_totp_enroll(monkeypatch, email=TEST_EMAIL) -> AppTest:
    # Same one-shot-gate limitation as _dlg_change_password — call the
    # @st.dialog function directly and unconditionally.
    script = f"""
from uvalu.pages_.settings import _dlg_totp_enroll
_dlg_totp_enroll({email!r})
"""
    at = AppTest.from_string(script, default_timeout=60)
    at.run()
    assert not at.exception, [str(e.value) for e in at.exception]
    return at


class TestDlgTotpEnroll:
    def test_shows_qr_and_code_input(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        at = _run_dlg_totp_enroll(monkeypatch)
        assert len(at.image) == 1
        assert any(w.key == "totp_enroll_code" for w in at.text_input)
        assert any(b.label == "Confirm and enable" for b in at.button)

    def test_correct_code_enables_totp(self, isolated_data, monkeypatch):
        import pyotp
        auth.register(TEST_EMAIL, "password123")
        at = _run_dlg_totp_enroll(monkeypatch)
        secret = auth._load_users()[TEST_EMAIL]["totp_secret"]
        at.text_input(key="totp_enroll_code").set_value(pyotp.TOTP(secret).now())
        confirm_btn = [b for b in at.button if b.label == "Confirm and enable"][0]
        confirm_btn.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert auth.is_totp_enabled(TEST_EMAIL)

    def test_wrong_code_shows_error(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        at = _run_dlg_totp_enroll(monkeypatch)
        at.text_input(key="totp_enroll_code").set_value("000000")
        confirm_btn = [b for b in at.button if b.label == "Confirm and enable"][0]
        confirm_btn.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert not auth.is_totp_enabled(TEST_EMAIL)
        assert "didn't match" in "".join(e.value for e in at.error)

    def test_cancel_does_not_enable_totp(self, isolated_data, monkeypatch):
        # A cancel-click's st.session_state.pop() inside this @st.dialog
        # function doesn't reliably surface via the outer at.session_state
        # after .click().run() (confirmed via a minimal repro — a plain
        # counter increment/pop showed the same staleness), so this checks
        # the actually-observable outcome (auth.py's own persisted state)
        # instead of asserting on at.session_state directly.
        auth.register(TEST_EMAIL, "password123")
        at = _run_dlg_totp_enroll(monkeypatch)
        cancel_btn = [b for b in at.button if b.label == "Cancel"][0]
        cancel_btn.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert not auth.is_totp_enabled(TEST_EMAIL)


def _run_with_backup_codes_flag(monkeypatch, codes) -> AppTest:
    script = f"""
import streamlit as st
st.session_state["totp_new_backup_codes"] = {codes!r}
from uvalu.pages_.settings import _dlg_backup_codes_shown
_dlg_backup_codes_shown()
"""
    at = AppTest.from_string(script, default_timeout=60)
    at.run()
    assert not at.exception, [str(e.value) for e in at.exception]
    return at


class TestDlgBackupCodesShown:
    def test_shows_all_codes(self, isolated_data, monkeypatch):
        at = _run_with_backup_codes_flag(monkeypatch, ["aaaa-1111", "bbbb-2222"])
        assert len(at.code) == 1
        assert "aaaa-1111" in at.code[0].value
        assert "bbbb-2222" in at.code[0].value

    def test_done_disabled_until_checked(self, isolated_data, monkeypatch):
        at = _run_with_backup_codes_flag(monkeypatch, ["aaaa-1111"])
        done_btn = [b for b in at.button if b.label == "Done"][0]
        assert done_btn.disabled

    def test_done_clickable_once_checked(self, isolated_data, monkeypatch):
        # See TestDlgTotpEnroll's cancel test for why this doesn't assert on
        # at.session_state directly across a dialog button click.
        at = _run_with_backup_codes_flag(monkeypatch, ["aaaa-1111"])
        at.checkbox(key="totp_backup_saved").set_value(True)
        at.run()
        done_btn = [b for b in at.button if b.label == "Done"][0]
        assert not done_btn.disabled
        done_btn.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]


class TestBackupCodesCard:
    def test_shows_remaining_count(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        _enable_totp()
        at = _run(monkeypatch)
        html = "".join(m.value for m in at.markdown)
        assert "10 of 10 unused. Regenerating replaces the old set." in html

    def test_regenerate_opens_confirm_dialog(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        _enable_totp()
        at = _run(monkeypatch)
        regen_btn = [b for b in at.button if b.label == "Regenerate"][0]
        regen_btn.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert any(b.key == "totp_regen_confirm" for b in at.button)

    def test_confirm_regenerates_codes(self, isolated_data, monkeypatch):
        # Same AppTest dialog-session_state limitation noted on TestDlgTotpEnroll's
        # cancel test -- checks the real backup-codes-remaining count (an
        # observable, persisted side effect) rather than the staged
        # at.session_state["totp_new_backup_codes"] value.
        auth.register(TEST_EMAIL, "password123")
        _enable_totp()
        at = _run(monkeypatch)
        regen_btn = [b for b in at.button if b.label == "Regenerate"][0]
        regen_btn.click().run()
        confirm_btn = [b for b in at.button if b.key == "totp_regen_confirm"][0]
        confirm_btn.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert auth.backup_codes_remaining(TEST_EMAIL) == 10


class TestTrustedDevicesCard:
    def test_shows_device_count(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        _enable_totp()
        auth.trust_this_device(TEST_EMAIL)
        at = _run(monkeypatch)
        html = "".join(m.value for m in at.markdown)
        assert "1 device" in html

    def test_revoke_all_clears_devices(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        _enable_totp()
        auth.trust_this_device(TEST_EMAIL)
        at = _run(monkeypatch)
        revoke_btn = [b for b in at.button if b.label == "Revoke all"][0]
        revoke_btn.click().run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert auth.trusted_device_count(TEST_EMAIL) == 0

    def test_revoke_all_disabled_when_no_devices(self, isolated_data, monkeypatch):
        auth.register(TEST_EMAIL, "password123")
        _enable_totp()
        at = _run(monkeypatch)
        revoke_btn = [b for b in at.button if b.label == "Revoke all"][0]
        assert revoke_btn.disabled


# ── Passkeys stub (Phase 3 — UI only, no WebAuthn) ──────────────────────────

class TestPasskeysStubRow:
    def test_says_not_available_in_beta_and_disables_manage(self, isolated_data, monkeypatch):
        at = _run(monkeypatch)
        html = "".join(m.value for m in at.markdown)
        assert "Passkeys" in html
        assert "Not available in the beta." in html
        assert "PHASE 3" not in html                               # no internal roadmap labels
        manage_btn = [b for b in at.button if b.label == "Manage"][0]
        assert manage_btn.disabled

    def test_shown_regardless_of_password_or_totp_state(self, isolated_data, monkeypatch):
        # No account registered at all (provider-only) -- the row still
        # appears, unlike the TOTP-gated rows above it.
        at = _run(monkeypatch)
        assert any(b.label == "Manage" for b in at.button)


def test_uses_the_supporting_page_frame(isolated_data, monkeypatch):
    """Proposal 1d/1e: 24px title in the shared frame, the section list as the
    rail, sentence-case card headers."""
    at = _run(monkeypatch)
    html = "".join(m.value for m in at.markdown)
    assert 'class="uv-support-title"' in html and ">Settings<" in html
    for anchor, _label in settings_page._SECTIONS:
        assert f'href="#{anchor}"' in html and f'id="{anchor}"' in html
    assert 'class="uv-card-head uv-set-head"' in html
    assert "text-transform:uppercase" not in html.split('class="uv-card-head')[1][:200]


def test_table_density_persists(isolated_data, monkeypatch):
    """G15: the Display card's Table density writes the density setting."""
    at = _run(monkeypatch)
    at.segmented_control(key="set_density_seg").set_value("compact").run()
    assert not at.exception, [str(e.value) for e in at.exception]
    assert settings.load_settings(TEST_EMAIL)["density"] == "compact"


def test_legal_card_is_legal_documents_with_english_only_note(isolated_data, monkeypatch):
    """G18."""
    at = _run(monkeypatch)
    html = "".join(m.value for m in at.markdown)
    assert "Legal documents" in html
    assert "During the beta these documents are available in English only." in html


def test_language_preview_band_holds_the_try_a_number_field(isolated_data, monkeypatch):
    """S09: "Preview" header (12px/500, --text), sentence-case labels, and the
    "Try a number" field inside the same --panel-2 band."""
    at = _run(monkeypatch)
    preview = next(m.value for m in at.markdown if "uv-i18n-preview" in m.value)
    assert "font-weight:500;color:var(--text);" in preview
    assert "text-transform:uppercase" not in preview
    field = at.text_input(key="set_i18n_try_input")
    field.set_value("1234").run()
    assert any("Saved as" in m.value for m in at.markdown)


# ── Privacy, Plan and Legal documents ────────────────────────────────────────

def test_privacy_plan_and_legal_cards(isolated_data, monkeypatch):
    at = _run(monkeypatch)
    html = "".join(m.value for m in at.markdown)
    for text in ("Export my data", "Delete account", "Privacy contact", "privacy&#64;uvalu.app",
                 "Free beta · by invitation", "Active", "Terms of use · beta", "Privacy notice · beta",
                 "No-personal-advice acknowledgement", "isn't recorded yet", "Not available",
                 "available in English only"):
        assert text in html, text
    assert "Export" in [b.label for b in at.download_button]


def _two_users():
    import auth
    auth.register("admin@example.com", "password123-long")
    auth.register("test@example.com", "password123-long")


def _run_delete_dialog():
    # A dialog opened by a button doesn't survive a second AppTest run (the
    # button is false then), so the dialog is called directly, every run.
    at = AppTest.from_string(USER_SETUP_SRC + """
from uvalu.pages_ import settings as settings_page
settings_page._dlg_delete_account("test@example.com")
""", default_timeout=60)
    at.run()
    assert not at.exception, [str(e.value) for e in at.exception]
    return at


def test_delete_account_needs_the_typed_email(isolated_data):
    import auth
    _two_users()
    portfolio.save_watchlist({"AAA.BR"})
    at = _run_delete_dialog()
    at.text_input(key="set_delete_email").set_value("someone@else.com").run()
    at.button(key="set_delete_confirm").click().run()
    assert "doesn't match" in "".join(e.value for e in at.error)
    assert "test@example.com" in auth._load_users()


def test_delete_account_removes_the_account_and_signs_out(isolated_data):
    import auth
    _two_users()
    portfolio.save_watchlist({"AAA.BR"})
    at = _run_delete_dialog()
    at.text_input(key="set_delete_email").set_value("Test@Example.com").run()
    at.button(key="set_delete_confirm").click().run()
    assert "test@example.com" not in auth._load_users()
    assert portfolio.load_watchlist() == set()
    assert at.query_params.get("logout") == ["1"]


def test_last_admin_cannot_delete(isolated_data, monkeypatch):
    import auth
    auth.register("test@example.com", "password123-long")      # first account: Admin
    at = _run(monkeypatch, role="Admin")
    at.button(key="privacy_delete").click().run()
    assert "only active Admin" in "".join(e.value for e in at.error)
    assert not [b for b in at.button if b.key == "set_delete_confirm"]
