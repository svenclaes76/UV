"""Unit tests for workbook.py — the user-facing export and import
(docs/import-export-plan.md, IE1–IE5)."""
from __future__ import annotations

import io
from datetime import date, datetime

import openpyxl
import pandas as pd
import pytest

import backup
import cash
import fx
import portfolio
import settings
import workbook

EMAIL = "test@example.com"


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv("ENCRYPTION_KEY", "unit-test-key-123")
    monkeypatch.setattr(portfolio, "_BASE_DIR", tmp_path / "portfolio")
    monkeypatch.setattr(settings, "_DATA_DIR", tmp_path / "settings_data")
    monkeypatch.setattr(settings, "_SHARED_FILE", tmp_path / "settings_data" / "shared.json")
    monkeypatch.setattr(backup, "_SHARED_FILE", tmp_path / "settings_data" / "shared.json")
    monkeypatch.setattr(backup, "_BACKUPS_DIR", tmp_path / "backups")
    monkeypatch.setattr(backup, "_BACKUPS_MANIFEST", tmp_path / "backups" / "manifest.json")
    (tmp_path / "settings_data").mkdir()
    portfolio.set_user(EMAIL)
    yield
    portfolio.set_user("")


@pytest.fixture
def ecb(monkeypatch):
    rates = {"USD": 0.85, "CHF": 1.07}

    def fake(ccy, base="EUR", on=None):
        ccy = ccy.upper()
        if ccy == base:
            return fx.FxQuote(1.0, on or date.today(), "base")
        if ccy not in rates:
            raise fx.FxUnavailable("no rate")
        return fx.FxQuote(rates[ccy], on or date.today(), "ecb")

    monkeypatch.setattr(fx, "get_rate", fake)
    monkeypatch.setattr(fx, "rates_frame", lambda currencies, start, end=None, base="EUR": pd.DataFrame(
        {c: [rates[c]] for c in currencies if c in rates}, index=pd.to_datetime(["2020-01-01"])))
    return rates


def _sample_data():
    """A small account: deposit, two buys, a partial sale, dividends,
    watchlist, custom ticker, targets and withholding."""
    cash.post_manual("Deposit", "2024-01-02", 10000, note="Opening transfer")
    portfolio.record_buy({"name": "AB InBev", "ticker": "ABI.BR", "shares": 40, "purchase_value": 2000.0,
                          "dividends": 0.0, "date_in": "2024-01-10T00:00:00"}, fee=7.5)
    portfolio.record_buy({"name": "Roche", "ticker": "ROG.SW", "shares": 8, "purchase_value": 2065.0,
                          "dividends": 0.0, "date_in": "2024-02-01T00:00:00"})
    portfolio.record_sell("ABI.BR", 10, 60.0, 3.0, "2024-09-01T00:00:00")
    portfolio.add_dividend({"name": "AB InBev", "ticker": "ABI.BR", "shares": 40, "amount": 32.8,
                            "amount_per_share": 0.82, "currency": "EUR", "tax_rate": 0.0, "tax_amount": 0.0,
                            "div_type": "Cash", "source": "manual",
                            "ex_date": "2024-05-08T00:00:00", "date": "2024-05-10T00:00:00"})
    portfolio.save_watchlist({"ASML.AS", "BDC"})
    portfolio.save_manual_tickers({"BDC": "Belden Inc"})
    portfolio.save_targets({"sectors": {"Technology": 0.3}, "tickers": {"ABI.BR": 0.1}, "hhi_max": 0.2})
    s = settings.load_settings(EMAIL)
    s["dividend_withholding"] = {**s["dividend_withholding"], "swiss": 35.0}
    settings.save_settings(s, EMAIL)


def _book(data: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(data))


def _sheet_rows(wb: openpyxl.Workbook, title: str) -> list[tuple]:
    return list(wb[title].iter_rows(values_only=True))


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── export ────────────────────────────────────────────────────────────────────

class TestExport:
    def test_empty_account_gets_headers_only(self):
        """The empty portfolio export is the template: headers only."""
        wb = _book(workbook.export_workbook(EMAIL))
        assert wb.sheetnames == ["Positions", "Sold", "Dividends", "Cash", "Watchlist", "_meta"]
        for key, title in zip(workbook.PORTFOLIO_SHEETS, wb.sheetnames):
            assert _sheet_rows(wb, title) == [workbook.COLUMNS[key]]

    def test_portfolio_export_holds_only_portfolio_sheets(self, ecb):
        _sample_data()
        wb = _book(workbook.export_workbook(EMAIL))
        assert wb.sheetnames == ["Positions", "Sold", "Dividends", "Cash", "Watchlist", "_meta"]
        assert wb["_meta"].sheet_state == "hidden"
        for key, title in zip(workbook.PORTFOLIO_SHEETS, wb.sheetnames):
            assert _sheet_rows(wb, title)[0] == workbook.COLUMNS[key]

    def test_full_export_holds_everything(self, ecb):
        _sample_data()
        wb = _book(workbook.export_all(EMAIL))
        assert wb.sheetnames == ["Account", "Settings", "Positions", "Sold", "Dividends", "Cash",
                                 "Watchlist", "Automatic cash", "_meta"]
        assert _sheet_rows(wb, "Watchlist")[1:] == [("ASML.AS", None), ("BDC", "Belden Inc")]

    def test_only_user_entered_values(self, ecb):
        _sample_data()
        wb = _book(workbook.export_workbook(EMAIL))
        pos = _sheet_rows(wb, "Positions")
        assert len(pos) == 3
        abi = next(r for r in pos[1:] if r[1] == "ABI.BR")
        assert abi[2:5] == (30, 1500.0, 7.5)
        assert isinstance(abi[5], datetime) and abi[5].date() == date(2024, 1, 10)
        sold = _sheet_rows(wb, "Sold")
        assert sold[1][1:6] == ("ABI.BR", 10, 500.0, 600.0, 3.0)
        cash_rows = _sheet_rows(wb, "Cash")
        assert cash_rows[1:] == [(datetime(2024, 1, 2), "Deposit", 10000, "EUR", None, "Opening transfer")]
        assert _sheet_rows(wb, "Watchlist")[1:] == [("ASML.AS", None), ("BDC", "Belden Inc")]

    def test_account_sheet_holds_no_secrets(self, monkeypatch):
        import auth
        monkeypatch.setattr(auth, "_load_users", lambda: {EMAIL: {
            "role": "Analyst", "status": "Active", "created_at": "2026-07-07T12:02:00",
            "password_hash": "$2b$secret", "totp_secret": "TOTPSECRET", "totp_enabled": True,
            "backup_codes": ["code1"], "reset_token": "tok", "invite_token": "tok2",
            "sessions": [{"sid": "s1", "created_at": "2026-10-01", "user_agent": "Firefox", "revoked": False}]}})
        wb = _book(workbook.export_all(EMAIL))
        rows = dict(r for r in _sheet_rows(wb, "Account")[1:] if r[0] != "session")
        assert rows["email"] == EMAIL and rows["two_step_verification"] == "yes"
        text = repr(_sheet_rows(wb, "Account"))
        for secret in ("$2b$secret", "TOTPSECRET", "code1", "tok", "s1"):
            assert secret not in text
        assert any(r[0] == "session" and "Firefox" in r[1] for r in _sheet_rows(wb, "Account"))

    def test_settings_sheet_lists_every_setting_by_section(self, ecb):
        _sample_data()
        portfolio.set_dividend_meta("ABI.BR", frequency="Semi-annual")
        rows = _sheet_rows(_book(workbook.export_all(EMAIL)), "Settings")
        assert rows[0] == ("section", "setting", "value")
        by = {(a, b): c for a, b, c in rows[1:]}
        assert by[("Language & region", "Language")] == "Not set (default)"
        assert by[("Display", "Table density")] == "comfortable"
        assert by[("Dividend withholding", "SIX Swiss Exchange")] == "35%"
        assert by[("Dividend withholding", "Euronext Brussels")] == "0%"
        assert by[("Target allocation", "Sector · Technology")] == "30%"
        assert by[("Target allocation", "Ticker · ABI.BR")] == "10%"
        assert by[("Target allocation", "Concentration ceiling (HHI)")] == "20%"
        assert by[("Dividends", "Payment frequency · ABI.BR")] == "Semi-annual"
        assert by[("Portfolio", "Base currency")] == "EUR"
        sections = [a for a, _b, _c in rows[1:]]
        assert sections == sorted(sections, key=sections.index)     # each section in one block

    def test_automatic_cash_sheet(self, ecb):
        _sample_data()
        auto = _sheet_rows(_book(workbook.export_all(EMAIL)), "Automatic cash")[1:]
        assert {r[1] for r in auto} == {"Buy", "Sell", "Dividend"}
        assert all(r[6] for r in auto)                   # linked_to
        assert "Deposit" not in {r[1] for r in auto}      # the typed deposit is on the Cash sheet

    def test_full_export_imports_only_the_portfolio(self, ecb):
        _sample_data()
        parsed = workbook.parse_workbook(workbook.export_all(EMAIL))
        assert parsed.issues == []
        assert parsed.sheets == ["positions", "sold", "dividends", "cash", "watchlist"]
        assert len(parsed.cash) == 1

    def test_formula_like_text_is_written_as_text(self):
        """SR7: a note starting with = + - @ must not become a formula in Excel,
        and must import back unchanged."""
        note = '=HYPERLINK("http://example.com","x")'
        cash.post_manual("Deposit", date(2024, 1, 2), 100, "EUR", note=note)
        data = workbook.export_workbook(EMAIL)
        cell = _book(data)["Cash"]["F2"]
        assert cell.value == note and cell.data_type == "s"
        parsed = workbook.parse_workbook(data)
        assert [r["note"] for r in parsed.cash] == [note]

    def test_dividend_without_currency_exports_exchange_currency(self):
        portfolio.save_div_hist(pd.DataFrame([{"ticker": "ROG.SW", "shares": 8, "amount": 76.8,
                                               "date": "2024-03-20T00:00:00"}]))
        wb = _book(workbook.export_workbook(EMAIL))
        assert _sheet_rows(wb, "Dividends")[1][4] == "CHF"

    def test_rows_added_to_an_export_import_in_add_mode(self, ecb):
        _sample_data()
        wb = _book(workbook.export_workbook(EMAIL))
        wb["Positions"].append(["ASML", "ASML.AS", 2, 1300, None, datetime(2025, 3, 3)])
        buf = io.BytesIO()
        wb.save(buf)
        plan = workbook.plan_import(workbook.parse_workbook(buf.getvalue()), "add")
        assert (plan.new["positions"], plan.skipped["positions"]) == (1, 2)


# ── parsing ───────────────────────────────────────────────────────────────────

class TestParse:
    def test_not_a_workbook(self):
        p = workbook.parse_workbook(b"not a zip")
        assert len(p.issues) == 1

    def test_unknown_sheets_only(self):
        p = workbook.parse_workbook(_xlsx({"beleggingen": [["foo", "bar"], [1, 2]]}))
        assert len(p.issues) == 1 and not p.sheets

    def test_sheets_found_by_headers_not_names(self):
        p = workbook.parse_workbook(_xlsx({
            "Posities": [["Ticker", "Shares", "Purchase value", "date_in"], ["abi.br", 10, 500, "2024-01-10"]],
            "Volglijst": [["Ticker"], ["ASML.AS"]],
        }))
        assert p.issues == []
        assert p.positions[0]["ticker"] == "ABI.BR" and p.positions[0]["date_in"] == date(2024, 1, 10)
        assert p.sheets == ["positions", "watchlist"] and p.watchlist == {"ASML.AS": ""}

    def test_text_numbers_and_dates(self):
        p = workbook.parse_workbook(_xlsx({"x": [
            ["ticker", "shares", "purchase_value", "fee", "date_in"],
            ["A.BR", "1.000", "1.974,36", "2,50", "25/11/2025"],
            ["B.BR", "10", "1,974.36", "", "2025-11-25T00:00:00"],
        ]}))
        assert p.issues == []
        assert p.positions[0]["purchase_value"] == pytest.approx(1974.36)
        assert p.positions[0]["fee"] == pytest.approx(2.5)
        assert p.positions[0]["date_in"] == date(2025, 11, 25)
        assert p.positions[1]["purchase_value"] == pytest.approx(1974.36)

    def test_reads_a_pre_format_export(self):
        """The export before this rework: extra derived columns, ISO text dates."""
        p = workbook.parse_workbook(_xlsx({
            "Positions": [["name", "google_ticker", "ticker", "shares", "purchase_value", "target_price",
                           "dividends", "date_in", "account", "purchase_price", "trade_id"],
                          ["Aedifica", "EBR:AED", "AED.BR", 30, 1974.36, 100, 84, "2025-11-25T00:00:00",
                           None, 65.81, "TRD-0001"]],
            "Dividends": [["name", "google_ticker", "ticker", "shares", "amount", "date", "declaration_date",
                           "ex_date", "record_date", "amount_per_share", "div_type", "source", "currency"],
                          ["Roche", "SWX:ROG", "ROG.SW", 8, 76.8, "2024-03-20T00:00:00", None,
                           "2024-03-20T00:00:00", None, 9.6, "Cash", "manual", None]],
        }))
        assert p.issues == []
        assert p.positions[0]["purchase_value"] == 1974.36
        assert p.dividends[0]["currency"] == "CHF"

    def test_row_issues_name_the_row(self):
        p = workbook.parse_workbook(_xlsx({"x": [
            ["ticker", "shares", "purchase_value", "date_in"],
            ["", 10, 500, "2024-01-10"],
            ["A.BR", "ten", 500, "2024-01-10"],
            ["B.BR", 10, 500, "2999-01-01"],
            ["C.BR", 10, 500, "2024-01-10"],
        ]}))
        assert sorted({i.row for i in p.issues}) == [2, 3, 4]
        assert [r["ticker"] for r in p.positions] == ["C.BR"]

    def test_cash_and_dividend_rules(self):
        p = workbook.parse_workbook(_xlsx({
            "c": [["date", "type", "amount", "currency"], ["2024-01-01", "Gift", 5, "EUR"],
                  ["2024-01-01", "deposit", 0, "EUR"], ["2024-01-01", "Adjustment", 0, None]],
            "d": [["ticker", "shares", "amount", "date", "ex_date", "div_type"],
                  ["A.BR", 10, 5, "2024-01-01", "2024-02-01", "Cash"],
                  ["B.BR", 10, 5, "2024-03-01", None, "bonus"],
                  ["C.BR", 10, 5, "2024-03-01", None, "special"]],
        }))
        assert sorted((i.sheet, i.row) for i in p.issues) == [("cash", 2), ("cash", 3), ("dividends", 2),
                                                              ("dividends", 3)]
        assert p.cash[0]["type"] == "Adjustment"
        assert p.dividends[0]["div_type"] == "Special" and p.dividends[0]["ex_date"] == date(2024, 3, 1)

    def test_amount_per_share_column(self):
        p = workbook.parse_workbook(_xlsx({"d": [["ticker", "shares", "amount_per_share", "date"],
                                                 ["A.BR", 40, 0.82, "2024-05-10"]]}))
        assert p.issues == [] and p.dividends[0]["amount"] == 32.8

    def test_newer_format_is_refused(self):
        p = workbook.parse_workbook(_xlsx({"Positions": [["ticker", "shares", "purchase_value", "date_in"]],
                                           "_meta": [["format_version", 99]]}))
        assert p.issues


# ── import ────────────────────────────────────────────────────────────────────

def _state():
    pf = portfolio.load_portfolio()
    sd = portfolio.load_sold()
    dv = portfolio.load_div_hist()
    return {
        "positions": sorted((r["ticker"], r["shares"], r["purchase_value"], r.get("fee"), r["date_in"][:10])
                            for r in pf.to_dict("records")),
        "sold": sorted((r["ticker"], r["shares"], r["purchase_value"], r["sale_value"], r["date_out"][:10])
                       for r in sd.to_dict("records")),
        "dividends": sorted((r["ticker"], r["amount"], r["currency"], r["date"][:10])
                            for r in dv.to_dict("records")),
        "balance": cash.balance(),
        "watchlist": portfolio.load_watchlist(),
        "custom": portfolio.load_manual_tickers(),
        "targets": portfolio.load_targets(),
        "withholding": settings.load_settings(EMAIL)["dividend_withholding"],
    }


class TestImport:
    def test_replace_restores_an_export(self, ecb):
        _sample_data()
        before = _state()
        data = workbook.export_workbook(EMAIL)
        # Change the portfolio, then restore.
        portfolio.record_buy({"name": "X", "ticker": "X.BR", "shares": 1, "purchase_value": 10.0,
                              "dividends": 0.0, "date_in": "2025-01-01T00:00:00"})
        parsed = workbook.parse_workbook(data)
        assert parsed.issues == []
        workbook.apply_import(parsed, "replace", EMAIL)
        assert _state() == before
        assert len(backup.list_backups()) == 1

    def test_replace_restores_watchlist_and_keeps_settings(self, ecb):
        _sample_data()
        data = workbook.export_workbook(EMAIL)
        portfolio.save_watchlist({"OTHER.PA"})
        portfolio.save_manual_tickers({"OTHER.PA": "Other"})
        workbook.apply_import(workbook.parse_workbook(data), "replace", EMAIL)
        assert portfolio.load_watchlist() == {"ASML.AS", "BDC"}
        assert portfolio.load_manual_tickers() == {"BDC": "Belden Inc"}
        assert portfolio.load_targets()["hhi_max"] == 0.2
        assert settings.load_settings(EMAIL)["dividend_withholding"]["swiss"] == 35.0

    def test_add_merges_the_watchlist(self):
        portfolio.save_watchlist({"A.BR"})
        data = _xlsx({"w": [["ticker", "name"], ["A.BR", None], ["XYZ", "Xyz Corp"], ["b.as", None]]})
        plan = workbook.apply_import(workbook.parse_workbook(data), "add", EMAIL)
        assert (plan.new["watchlist"], plan.skipped["watchlist"]) == (2, 1)
        assert portfolio.load_watchlist() == {"A.BR", "XYZ", "B.AS"}
        assert portfolio.load_manual_tickers() == {"XYZ": "Xyz Corp"}

    def test_replace_links_cash_to_new_trade_ids(self, ecb):
        _sample_data()
        workbook.apply_import(workbook.parse_workbook(workbook.export_workbook(EMAIL)), "replace", EMAIL)
        ids = set(portfolio.load_portfolio()["trade_id"]) | set(portfolio.load_sold()["trade_id"])
        linked = {e["ref_id"] for e in cash.load_ledger() if e.get("ref_kind") == "trade"}
        assert linked == ids

    def test_add_skips_what_is_already_there(self, ecb):
        _sample_data()
        before = _state()
        parsed = workbook.parse_workbook(workbook.export_workbook(EMAIL))
        plan = workbook.plan_import(parsed, "add")
        assert plan.skipped == {"positions": 2, "sold": 1, "dividends": 1, "cash": 1, "watchlist": 2}
        assert plan.new == {"positions": 0, "sold": 0, "dividends": 0, "cash": 0, "watchlist": 0}
        workbook.apply_import(parsed, "add", EMAIL)
        assert _state() == before

    def test_add_rows_post_cash_with_top_up(self):
        data = _xlsx({"Positions": [list(workbook.COLUMNS["positions"]),
                                    ["", "ABI.BR", 10, 500, 2.5, "2024-01-10"]]})
        workbook.apply_import(workbook.parse_workbook(data), "add", EMAIL)
        pos = portfolio.load_portfolio().iloc[0]
        assert (pos["name"], pos["purchase_price"], pos["fee"]) == ("ABI.BR", 50.0, 2.5)
        entries = cash.load_ledger()
        assert [e["type"] for e in entries] == ["Deposit", "Buy"]
        assert entries[0]["topup"] and cash.balance() == 0.0

    def test_sold_dividends_are_recomputed_by_share_count(self, ecb):
        data = _xlsx({
            "s": [["ticker", "shares", "purchase_value", "sale_value", "date_in", "date_out"],
                  ["ABI.BR", 30, 1500, 1800, "2024-01-10", "2024-12-01"]],
            "p": [["ticker", "shares", "purchase_value", "date_in"], ["ABI.BR", 10, 500, "2024-01-10"]],
            "d": [["ticker", "shares", "amount", "date"], ["ABI.BR", 40, 100, "2024-05-10"],
                  ["ABI.BR", 10, 50, "2025-05-10"]],
        })
        workbook.apply_import(workbook.parse_workbook(data), "replace", EMAIL)
        sold = portfolio.load_sold().iloc[0]
        # 100 gross, 30% Belgian withholding = 70 net; the sold lot held 30 of 40 shares.
        assert sold["dividends"] == pytest.approx(52.5)
        assert sold["annual_return_pct"] > 0

    def test_missing_fx_rate_stops_before_writing(self, ecb):
        _sample_data()
        before = _state()
        data = _xlsx({"c": [["date", "type", "amount", "currency"], ["2024-01-01", "Deposit", 50, "JPY"]]})
        with pytest.raises(workbook.ImportError_):
            workbook.apply_import(workbook.parse_workbook(data), "replace", EMAIL)
        assert _state() == before

    def test_file_with_issues_is_refused(self):
        p = workbook.parse_workbook(_xlsx({"x": [["ticker", "shares", "purchase_value", "date_in"],
                                                 ["A.BR", -1, 500, "2024-01-10"]]}))
        with pytest.raises(workbook.ImportError_):
            workbook.apply_import(p, "add", EMAIL)

    def test_foreign_cash_with_given_rate(self):
        data = _xlsx({"c": [["date", "type", "amount", "currency", "fx_rate"],
                            ["2024-01-01", "Deposit", 100, "USD", 0.9]]})
        workbook.apply_import(workbook.parse_workbook(data), "add", EMAIL)
        e = cash.load_ledger()[0]
        assert (e["currency"], e["fx_rate"], e["amount_base"]) == ("USD", 0.9, 90.0)
